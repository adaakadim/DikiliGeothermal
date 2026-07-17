import os
import csv
from datetime import datetime
from typing import Dict, Any, List

try:
    from fpdf import FPDF
except ImportError:
    FPDF = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
except ImportError:
    Workbook = None

class EnterpriseReportGenerator:
    def __init__(self, output_dir: str = "data/reports"):
        self.output_dir = output_dir
        if not os.path.exists(self.output_dir):
            os.makedirs(self.output_dir)
            
    def generate_report(
        self, 
        report_type: str,
        export_format: str,
        operator_name: str,
        scenario_name: str,
        date_range_start: str,
        date_range_end: str,
        data: Dict[str, Any]
    ) -> str:
        
        timestamp = datetime.now()
        ts_str = timestamp.strftime("%Y%m%d_%H%M%S")
        filename_base = f"{report_type.replace(' ', '_')}_{ts_str}"
        
        # Prepare Metadata Header
        meta = {
            "Project Name": "Dikili Geothermal Platform",
            "Plant Name": "Dikili Main Geothermal Plant",
            "Report Type": report_type,
            "Date": timestamp.strftime("%Y-%m-%d %H:%M:%S"),
            "Operator": operator_name,
            "Scenario": scenario_name,
            "Version": "1.0.0 Release Candidate",
            "Period Start": date_range_start,
            "Period End": date_range_end
        }
        
        if export_format == "PDF":
            return self._generate_pdf(filename_base, meta, data)
        elif export_format == "Excel (.xlsx)":
            return self._generate_excel(filename_base, meta, data)
        elif export_format == "CSV":
            return self._generate_csv(filename_base, meta, data)
        else:
            raise ValueError(f"Unsupported format: {export_format}")
            
    def _generate_pdf(self, filename_base: str, meta: Dict[str, str], data: Dict[str, Any]) -> str:
        if not FPDF:
            raise RuntimeError("fpdf2 library is not installed.")
            
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", style="B", size=16)
        
        # Title
        pdf.cell(200, 10, txt=f"Dikili Geothermal - {meta['Report Type']}", ln=True, align='C')
        pdf.ln(10)
        
        # Metadata
        pdf.set_font("Helvetica", size=10)
        for k, v in meta.items():
            pdf.cell(200, 6, txt=f"{k}: {v}", ln=True)
        pdf.ln(10)
        
        # Sections
        sections = [
            ("Plant Overview", "overview"),
            ("Production Summary", "production"),
            ("District Heating Summary", "heating"),
            ("Energy Summary", "energy"),
            ("Equipment Status", "equipment"),
            ("Alarm Summary", "alarms"),
            ("Maintenance Summary", "maintenance"),
            ("AI Summary", "ai"),
            ("Engineering Recommendations", "recommendations")
        ]
        
        for title, key in sections:
            pdf.set_font("Helvetica", style="B", size=12)
            pdf.cell(200, 8, txt=title, ln=True)
            pdf.set_font("Helvetica", size=10)
            
            section_data = data.get(key, {})
            if isinstance(section_data, dict):
                for k, v in section_data.items():
                    pdf.multi_cell(0, 6, txt=f"{k}: {v}")
            elif isinstance(section_data, list):
                for item in section_data:
                    pdf.multi_cell(0, 6, txt=f"- {item}")
            else:
                pdf.multi_cell(0, 6, txt=str(section_data))
            pdf.ln(5)
            
        filepath = os.path.join(self.output_dir, f"{filename_base}.pdf")
        pdf.output(filepath)
        return filepath

    def _generate_excel(self, filename_base: str, meta: Dict[str, str], data: Dict[str, Any]) -> str:
        if not Workbook:
            raise RuntimeError("openpyxl library is not installed.")
            
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        
        bold_font = Font(bold=True)
        header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
        
        # Write Metadata
        ws.append(["Report Metadata"])
        ws.cell(row=1, column=1).font = bold_font
        for k, v in meta.items():
            ws.append([k, str(v)])
            
        ws.append([])
        
        sections = [
            ("Plant Overview", "overview"),
            ("Production Summary", "production"),
            ("District Heating Summary", "heating"),
            ("Energy Summary", "energy"),
            ("Equipment Status", "equipment"),
            ("Alarm Summary", "alarms"),
            ("Maintenance Summary", "maintenance"),
            ("AI Summary", "ai"),
            ("Engineering Recommendations", "recommendations")
        ]
        
        row_idx = ws.max_row + 1
        for title, key in sections:
            ws.cell(row=row_idx, column=1, value=title).font = bold_font
            ws.cell(row=row_idx, column=1).fill = header_fill
            row_idx += 1
            
            section_data = data.get(key, {})
            if isinstance(section_data, dict):
                for k, v in section_data.items():
                    ws.cell(row=row_idx, column=1, value=k)
                    ws.cell(row=row_idx, column=2, value=str(v))
                    row_idx += 1
            elif isinstance(section_data, list):
                for item in section_data:
                    ws.cell(row=row_idx, column=1, value=str(item))
                    row_idx += 1
            else:
                ws.cell(row=row_idx, column=1, value=str(section_data))
                row_idx += 1
            row_idx += 1
            
        filepath = os.path.join(self.output_dir, f"{filename_base}.xlsx")
        wb.save(filepath)
        return filepath

    def _generate_csv(self, filename_base: str, meta: Dict[str, str], data: Dict[str, Any]) -> str:
        filepath = os.path.join(self.output_dir, f"{filename_base}.csv")
        with open(filepath, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            
            # Metadata
            writer.writerow(["--- METADATA ---"])
            for k, v in meta.items():
                writer.writerow([k, v])
            writer.writerow([])
            
            # Data Sections
            sections = [
                ("Plant Overview", "overview"),
                ("Production Summary", "production"),
                ("District Heating Summary", "heating"),
                ("Energy Summary", "energy"),
                ("Equipment Status", "equipment"),
                ("Alarm Summary", "alarms"),
                ("Maintenance Summary", "maintenance"),
                ("AI Summary", "ai"),
                ("Engineering Recommendations", "recommendations")
            ]
            
            for title, key in sections:
                writer.writerow([f"--- {title.upper()} ---"])
                section_data = data.get(key, {})
                if isinstance(section_data, dict):
                    for k, v in section_data.items():
                        writer.writerow([k, v])
                elif isinstance(section_data, list):
                    for item in section_data:
                        writer.writerow([item])
                else:
                    writer.writerow([section_data])
                writer.writerow([])
                
        return filepath
